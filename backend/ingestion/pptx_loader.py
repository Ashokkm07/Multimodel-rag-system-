from pptx import Presentation


def load_pptx(pptx_path):
    """
    Extract text from a PowerPoint presentation.
    Each slide becomes a document with slide metadata.
    """

    presentation = Presentation(pptx_path)

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        combined_text = "\n".join(slide_text).strip()

        if combined_text:
            documents.append({
                "text": combined_text,
                "metadata": {
                    "source": str(pptx_path),
                    "slide": slide_number
                }
            })

    return documents